"""#492 - question attachments: files and folders resolved into model-visible text.

The resolver is the single place that turns the local filesystem paths a
question carries into delimited attachment blocks. It routes every file through
the Knowledge Base's ``DocumentReader`` (no second parser), walks a dropped
directory, caps the total text per question, and reports per-file problems
instead of failing the turn.

Fixtures are built at test time in ``tmp_path`` -- the convention already used
by ``test_ingestion_reader`` -- so no binary blobs are checked into the tree.
"""

import pytest

from tests.test_ingestion_reader import _minimal_pdf

from src.utils.attachment_utils import (
    ATTACHMENT_CHAR_BUDGET,
    MAX_ATTACHMENT_FILES,
    TRUNCATION_MARKER,
    build_attachment_notice,
    resolve_attachments,
)

pytestmark = pytest.mark.unit


# ---------------------------------------------------------------------------
# Fixture builders - real files, real extractors, no mocks.
# ---------------------------------------------------------------------------


def _write_txt(directory, name="notes.txt", body="Quarterly notes for the team."):
    path = directory / name
    path.write_text(body, encoding="utf-8")
    return path


def _write_md(directory, name="readme.md"):
    path = directory / name
    path.write_text("# Title\n\nSome markdown body.\n", encoding="utf-8")
    return path


def _write_csv(directory, name="rows.csv"):
    path = directory / name
    path.write_text("quarter,revenue\nQ1,1200\nQ2,1400\n", encoding="utf-8")
    return path


def _write_pdf(directory, name="report.pdf"):
    path = directory / name
    path.write_bytes(_minimal_pdf(["Revenue grew steadily during the first quarter."]))
    return path


def _write_docx(directory, name="memo.docx"):
    import docx

    path = directory / name
    document = docx.Document()
    document.add_paragraph("Memo body for the whole company.")
    document.save(str(path))
    return path


def _write_xlsx(directory, name="budget.xlsx"):
    import openpyxl

    path = directory / name
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "Quarter"
    sheet["B1"] = "Revenue"
    sheet["A2"] = "Q1"
    sheet["B2"] = 1200
    workbook.save(str(path))
    return path


def _write_png(directory, name="scan.png"):
    path = directory / name
    path.write_bytes(bytes([137, 80, 78, 71, 13, 10, 26, 10]))
    return path


# ---------------------------------------------------------------------------
# Happy path, one file per supported extension.
# ---------------------------------------------------------------------------


class TestSupportedExtensions:
    @pytest.mark.parametrize(
        "builder, expected",
        [
            (_write_txt, "Quarterly notes"),
            (_write_md, "Some markdown body"),
            (_write_csv, "revenue"),
            (_write_pdf, "Revenue grew steadily"),
            (_write_docx, "Memo body"),
            (_write_xlsx, "Quarter"),
        ],
    )
    def test_each_supported_extension_yields_a_delimited_block(self, tmp_path, builder, expected):
        path = builder(tmp_path)

        resolved = resolve_attachments([str(path)])

        assert resolved.issues == []
        assert f"[Attached file: {path.name}]" in resolved.block
        assert "[End of attached file]" in resolved.block
        assert expected in resolved.block
        assert resolved.paths == [str(path)]
        assert resolved.truncated_names == []

    def test_no_attachments_yields_an_empty_result(self):
        for empty in (None, [], ["", "   "]):
            resolved = resolve_attachments(empty)
            assert resolved.block == ""
            assert resolved.paths == []
            assert resolved.issues == []

    def test_several_files_keep_the_order_they_were_attached_in(self, tmp_path):
        first = _write_txt(tmp_path, "first.txt", "First document body.")
        second = _write_txt(tmp_path, "second.txt", "Second document body.")

        resolved = resolve_attachments([str(first), str(second)])

        assert resolved.block.index("first.txt") < resolved.block.index("second.txt")
        assert resolved.paths == [str(first), str(second)]

    def test_the_same_file_attached_twice_is_read_once(self, tmp_path):
        path = _write_txt(tmp_path)

        resolved = resolve_attachments([str(path), str(path)])

        assert resolved.paths == [str(path)]
        assert resolved.block.count("[Attached file: notes.txt]") == 1


# ---------------------------------------------------------------------------
# Directory walk.
# ---------------------------------------------------------------------------


class TestDirectoryWalk:
    def test_a_directory_contributes_its_supported_files(self, tmp_path):
        folder = tmp_path / "dossier"
        folder.mkdir()
        _write_txt(folder, "a.txt", "Body of A.")
        _write_md(folder, "b.md")
        (folder / "ignored.odt").write_text("not supported", encoding="utf-8")

        resolved = resolve_attachments([str(folder)])

        assert "[Attached file: a.txt]" in resolved.block
        assert "[Attached file: b.md]" in resolved.block
        # An unsupported sibling inside a folder is skipped silently: the user
        # dropped a folder, not that file.
        assert "ignored.odt" not in resolved.block
        assert resolved.issues == []

    def test_the_walk_reaches_one_level_below_the_dropped_folder(self, tmp_path):
        folder = tmp_path / "dossier"
        (folder / "sub").mkdir(parents=True)
        _write_txt(folder, "top.txt", "Top level body.")
        _write_txt(folder / "sub", "nested.txt", "Nested body.")
        deeper = folder / "sub" / "deeper"
        deeper.mkdir()
        _write_txt(deeper, "too_deep.txt", "Too deep body.")

        resolved = resolve_attachments([str(folder)])

        assert "top.txt" in resolved.block
        assert "nested.txt" in resolved.block
        assert "too_deep.txt" not in resolved.block

    def test_hidden_entries_are_skipped(self, tmp_path):
        folder = tmp_path / "dossier"
        folder.mkdir()
        _write_txt(folder, ".secret.txt", "Hidden body.")
        _write_txt(folder, "visible.txt", "Visible body.")

        resolved = resolve_attachments([str(folder)])

        assert "visible.txt" in resolved.block
        assert ".secret.txt" not in resolved.block

    def test_an_empty_folder_is_reported_and_reads_nothing(self, tmp_path):
        folder = tmp_path / "vide"
        folder.mkdir()

        resolved = resolve_attachments([str(folder)])

        assert resolved.block == ""
        assert [issue.reason for issue in resolved.issues] == ["empty_folder"]

    def test_the_expanded_file_count_is_capped(self, tmp_path):
        folder = tmp_path / "dossier"
        folder.mkdir()
        for index in range(MAX_ATTACHMENT_FILES + 5):
            _write_txt(folder, f"file_{index:03d}.txt", f"Body {index}.")

        resolved = resolve_attachments([str(folder)])

        assert len(resolved.paths) == MAX_ATTACHMENT_FILES
        assert any(issue.reason == "too_many" for issue in resolved.issues)


# ---------------------------------------------------------------------------
# Per-file failures never fail the turn.
# ---------------------------------------------------------------------------


class TestPerFileFailures:
    def test_an_unsupported_extension_is_reported_not_raised(self, tmp_path):
        path = tmp_path / "archive.zip"
        path.write_bytes(b"PK\x03\x04")

        resolved = resolve_attachments([str(path)])

        assert [issue.reason for issue in resolved.issues] == ["unsupported"]
        assert resolved.paths == [str(path)]
        # The model is told the file exists and could not be read.
        assert "[Attached file: archive.zip]" in resolved.block
        assert "could not be read" in resolved.block

    def test_a_missing_file_is_reported(self, tmp_path):
        missing = tmp_path / "ghost.txt"

        resolved = resolve_attachments([str(missing)])

        assert [issue.reason for issue in resolved.issues] == ["missing"]
        assert "ghost.txt" in resolved.block

    def test_an_image_comes_back_as_pending_vision(self, tmp_path):
        path = _write_png(tmp_path)

        resolved = resolve_attachments([str(path)])

        assert [issue.reason for issue in resolved.issues] == ["pending_vision"]
        assert "scan.png" in resolved.block
        assert "could not be read" in resolved.block

    def test_a_scanned_pdf_comes_back_as_pending_vision(self, tmp_path):
        from pypdf import PdfWriter

        path = tmp_path / "scanned.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with path.open("wb") as handle:
            writer.write(handle)

        resolved = resolve_attachments([str(path)])

        assert [issue.reason for issue in resolved.issues] == ["pending_vision"]

    def test_an_empty_file_is_reported(self, tmp_path):
        path = tmp_path / "blank.txt"
        path.write_text("   \n\n", encoding="utf-8")

        resolved = resolve_attachments([str(path)])

        assert [issue.reason for issue in resolved.issues] == ["empty"]

    def test_an_extraction_error_is_reported_not_raised(self, tmp_path, monkeypatch):
        path = _write_txt(tmp_path)

        def _boom(self, _path):
            raise RuntimeError("extractor exploded")

        monkeypatch.setattr("src.ingestion.reader.DocumentReader.read", _boom)

        resolved = resolve_attachments([str(path)])

        assert [issue.reason for issue in resolved.issues] == ["unreadable"]
        assert "could not be read" in resolved.block

    def test_one_bad_file_does_not_stop_the_good_one(self, tmp_path):
        good = _write_txt(tmp_path, "good.txt", "Readable body here.")
        bad = tmp_path / "bad.zip"
        bad.write_bytes(b"PK\x03\x04")

        resolved = resolve_attachments([str(bad), str(good)])

        assert "Readable body here." in resolved.block
        assert [issue.reason for issue in resolved.issues] == ["unsupported"]


# ---------------------------------------------------------------------------
# Token budget.
# ---------------------------------------------------------------------------


class TestBudget:
    def test_content_over_the_budget_is_truncated_with_an_explicit_marker(self, tmp_path):
        body = "A" * (ATTACHMENT_CHAR_BUDGET + 5_000)
        path = _write_txt(tmp_path, "huge.txt", body)

        resolved = resolve_attachments([str(path)])

        assert TRUNCATION_MARKER in resolved.block
        assert resolved.truncated_names == ["huge.txt"]
        assert len(resolved.block) < len(body)

    def test_the_budget_spans_the_whole_question_not_one_file(self, tmp_path):
        body = "B" * (ATTACHMENT_CHAR_BUDGET - 100)
        first = _write_txt(tmp_path, "first.txt", body)
        second = _write_txt(tmp_path, "second.txt", "C" * 10_000)

        resolved = resolve_attachments([str(first), str(second)])

        assert resolved.truncated_names == ["second.txt"]
        assert resolved.block.count(TRUNCATION_MARKER) == 1

    def test_content_under_the_budget_is_untouched(self, tmp_path):
        path = _write_txt(tmp_path, "small.txt", "Short body.")

        resolved = resolve_attachments([str(path)])

        assert TRUNCATION_MARKER not in resolved.block
        assert resolved.truncated_names == []


# ---------------------------------------------------------------------------
# The notice the UI shows.
# ---------------------------------------------------------------------------


class TestNotice:
    def test_a_clean_resolution_produces_no_notice(self, tmp_path):
        resolved = resolve_attachments([str(_write_txt(tmp_path))])
        assert build_attachment_notice(resolved) == ""

    def test_failures_and_truncation_are_named_in_the_notice(self, tmp_path):
        bad = tmp_path / "archive.zip"
        bad.write_bytes(b"PK\x03\x04")
        huge = _write_txt(tmp_path, "huge.txt", "A" * (ATTACHMENT_CHAR_BUDGET + 1_000))

        resolved = resolve_attachments([str(bad), str(huge)])
        notice = build_attachment_notice(resolved)

        assert "archive.zip" in notice
        assert "huge.txt" in notice
        assert notice.endswith("\n\n")
